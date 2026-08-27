# apps\reports\services\export.py sirve para:
# generar y devolver informes del dashboard en formatos Excel y PDF

# Librerías estándar de Python
from datetime import date
from pathlib import Path
# Librerías externas (Django, requests, openpyxl, reportlab, etc.)
from django.conf import settings
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
# Imports internos de proyecto/apps
from .dashboard import build_report_dashboard


# HELPER PRIVADO: ajusta automáticamente el ancho de las columnas de una hoja Excel
def autofit(ws):

  for column in ws.columns:

    length = max(
      len(str(cell.value or ""))
      for cell in column
    )

    ws.column_dimensions[
      column[0].column_letter
    ].width = length + 3


# HELPER PRIVADO: aplica el formato y la alineación a las tablas del informe PDF
def apply_table_style(table, rows, font_size=8, left_columns=1):

  style = [
    ("BACKGROUND", (0,0), (-1,0), colors.lightgrey),
    ("TEXTCOLOR", (0,0), (-1,0), colors.black),
    ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
    ("GRID", (0,0), (-1,-1), 0.25, colors.grey),
    ("FONTSIZE", (0,0), (-1,-1), font_size),
    ("BOTTOMPADDING", (0,0), (-1,0), 5),
    ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
  ]

  # Cabeceras alineadas según columnas
  if left_columns > 0:
    style.append(
      ("ALIGN", (0,0), (left_columns-1,0), "LEFT")
    )

  style.append(
    ("ALIGN", (left_columns,0), (-1,0), "CENTER")
  )

  # Datos alineados según columnas
  if left_columns > 0:
    style.append(
      ("ALIGN", (0,1), (left_columns-1,-1), "LEFT")
    )

  style.append(
    ("ALIGN", (left_columns,1), (-1,-1), "CENTER")
  )

  # Filas alternas
  for row in range(1, rows):
    if row % 2 == 0:
      style.append(
        ("BACKGROUND", (0,row), (-1,row), colors.whitesmoke)
      )

  table.setStyle(TableStyle(style))


# HELPER PRIVADO: dibuja la cabecera del informe PDF con el logotipo
def draw_header(canvas, doc):

    logo = Path(settings.BASE_DIR) / "static" / "img" / "logo.png"

    canvas.saveState()

    canvas.drawImage(
      ImageReader(str(logo)),
      x=doc.pagesize[0] - 4.6*cm,
      y=doc.pagesize[1] - 2.5*cm,
      width=4.5*cm,
      height=1.5*cm,
      preserveAspectRatio=True,
      mask="auto",
    )

    canvas.restoreState()


# HELPER PRIVADO: añade el pie de página con fecha y numeración total de páginas
class NumberedCanvas(canvas.Canvas):

  def __init__(self, *args, **kwargs):
    super().__init__(*args, **kwargs)
    self._saved_page_states = []

  def showPage(self):
    self._saved_page_states.append(dict(self.__dict__))
    self._startPage()

  def save(self):
    total_pages = len(self._saved_page_states)

    for state in self._saved_page_states:
      self.__dict__.update(state)
      self.draw_footer(total_pages)
      super().showPage()

    super().save()

  def draw_footer(self, total_pages):

    self.saveState()

    # Línea superior
    self.setStrokeColor(colors.lightgrey)
    self.line(
      0.8 * cm,
      0.9 * cm,
      landscape(A4)[0] - 0.8 * cm,
      0.9 * cm,
    )

    self.setFont("Helvetica", 8)

    # Izquierda
    self.drawString(
      0.8 * cm,
      0.45 * cm,
      f'Informe generado automáticamente el {date.today().strftime("%d/%m/%Y")}',
    )

    # Derecha
    self.drawRightString(
      landscape(A4)[0] - 0.8 * cm,
      0.45 * cm,
      f'Página {self._pageNumber} de {total_pages}',
    )

    self.restoreState()


# ARCHIVO EXCEL: genera y devuelve el informe del dashboard en formato Excel
def export_xls():

  data = build_report_dashboard()

  wb = Workbook()

  # Resumen
  ws = wb.active
  ws.title = "Resumen"

  ws.append([
    "Indicador",
    "Valor",
  ])

  ws.append([
    "Proyectos",
    data["total_projects"],
  ])

  ws.append([
    "Tareas",
    data["total_tasks"],
  ])

  ws.append([
    "Completado (%)",
    data["completed_percentage"],
  ])

  ws.append([
    "Usuarios",
    data["total_users"],
  ])

  # Estado proyectos
  ws_projects = wb.create_sheet("Estado proyectos")

  ws_projects.append([
    "Código",
    "Proyecto",
    "Estado",
    "Equipo",
    "Tareas",
    "Asignadas (%)",
    "Pendientes",
    "Finalizadas",
    "Canceladas",
    "Progreso (%)",
    "Salud (%)",
  ])

  for project in data["project_statistics"]:

    ws_projects.append([
      project["code"],
      project["name"],
      "Archivado" if project["is_archived"] else "Activo",
      project["users"],
      project["total_tasks"],
      project["assignment_percentage"],
      project["pending_tasks"],
      project["completed_tasks"],
      project["cancelled"],
      project["completion_percentage"],
      project["health"],
    ])

  # Estado plantilla trabajo
  ws_team = wb.create_sheet("Estado plantilla")

  ws_team.append([
    "Usuario",
    "Rol",
    "Archivadas",
    "Activas",
    "Balance (%)",
    "Finalizadas",
    "Canceladas",
    "En plazo",
    "Vencidas",
    "Participación (%)",
    "Productividad (%)",
  ])

  for user in data["team_statistics"]:

    ws_team.append([
      user["username"],
      user["role"],
      user["archived_tasks"],
      user["active_tasks"],
      user["workload"],
      user["completed_tasks"],
      user["cancelled_tasks"],
      user["on_time_tasks"],
      user["overdue_tasks"],
      user["participation"],
      user["productivity"],
    ])

  header_fill = PatternFill(
    fill_type="solid",
    fgColor="DDDDDD",
  )

  header_font = Font(
    bold=True,
  )

  for sheet in wb.worksheets:

    for cell in sheet[1]:

      cell.fill = header_fill
      cell.font = header_font
      cell.alignment = Alignment(horizontal="center")

    autofit(sheet)

  response = HttpResponse(
    content_type=(
      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
  )

  filename = (
    f'projectflow_report_{date.today()}.xlsx'
  )

  response["Content-Disposition"] = (
    f'attachment; filename="{filename}"'
  )

  wb.save(response)

  return response


# ARCHIVO PDF: genera y devuelve el informe del dashboard en formato PDF
def export_pdf():

  data = build_report_dashboard()

  response = HttpResponse(
    content_type="application/pdf"
  )

  filename = f'projectflow_report_{date.today()}.pdf'

  response["Content-Disposition"] = (
    f'attachment; filename="{filename}"'
  )

  doc = SimpleDocTemplate(
    response,
    pagesize=landscape(A4),
    rightMargin=0.8 * cm,
    leftMargin=0.8 * cm,
    topMargin=2.2 * cm,
    bottomMargin=1.2 * cm,
  )

  styles = getSampleStyleSheet()

  story = []

  story.append(
    Paragraph(
      "<b>Informe ejecutivo</b>",
      styles["Title"],
    )
  )

  story.append(
    Paragraph(
      "Resumen indicadores",
      styles["Heading2"],
    )
  )

  story.append(Spacer(1, 0.2 * cm))

  summary = [
    ["Indicador", "Valor"],
    ["Proyectos", data["total_projects"]],
    ["Tareas", data["total_tasks"]],
    ["Usuarios", data["total_users"]],
    [
      "% completadas",
      f'{data["completed_percentage"]}%'
    ],
  ]

  summary_table = Table(
    summary,
    colWidths=[6 * cm, 3 * cm],
    hAlign="LEFT",
  )

  apply_table_style(summary_table, len(summary), font_size=11, left_columns=1)

  story.append(summary_table)

  story.append(Spacer(1, 0.8 * cm))

  story.append(PageBreak())

  story.append(
    Paragraph(
      "<b>Estado proyectos</b>",
      styles["Heading2"],
    )
  )

  story.append(Spacer(1, 0.2 * cm))

  project_rows = [[
    "Código",
    "Proyecto",
    "Estado",
    "Equipo",
    "Tareas",
    "Asignadas",
    "Pendientes",
    "Finalizadas",
    "Canceladas",
    "Progreso",
    "Salud",
  ]]

  for project in data["project_statistics"]:

    project_rows.append([
      project["code"],
      project["name"],
      "Archivado" if project["is_archived"] else "Activo",
      project["users"],
      project["total_tasks"],
      f'{project["assignment_percentage"]}%',
      project["pending_tasks"],
      project["completed_tasks"],
      project["cancelled"],
      f'{project["completion_percentage"]}%',
      f'{project["health"]}%',
    ])

  # Total: 27.6 cm
  project_table = Table(
    project_rows,
    colWidths=[
      1.9*cm,   # Código
      8*cm,     # Proyecto
      2.5*cm,   # Estado
      1.6*cm,   # Equipo
      1.6*cm,   # Tareas
      2*cm,     # Asignadas
      2*cm,     # Pendientes
      2*cm,     # Finalizadas
      2*cm,     # Canceladas
      2*cm,     # Progreso
      2*cm,     # Salud
    ],
    repeatRows=1,
  )

  apply_table_style(project_table, len(project_rows), left_columns=2)

  story.append(project_table)

  story.append(Spacer(1, 0.8 * cm))

  story.append(PageBreak())

  story.append(
    Paragraph(
      "<b>Estado plantilla de trabajo</b>",
      styles["Heading2"],
    )
  )

  story.append(Spacer(1, 0.2 * cm))

  team_rows = [[
    "Usuario",
    "Rol",
    "Archivadas",
    "Activas",
    "Balance",
    "Finalizadas",
    "Canceladas",
    "En plazo",
    "Vencidas",
    "Participación",
    "Productividad",
  ]]

  for user in data["team_statistics"]:

    team_rows.append([
      user["username"],
      user["role"],
      user["archived_tasks"],
      user["active_tasks"],
      f'{user["workload"]}%',
      user["completed_tasks"],
      user["cancelled_tasks"],
      user["on_time_tasks"],
      user["overdue_tasks"],
      f'{user["participation"]}%',
      f'{user["productivity"]}%',
    ])


  # Total: 27.6 cm
  team_table = Table(
    team_rows,
    colWidths=[
      4.8*cm,  # Usuario
      3*cm,    # Rol
      2.2*cm,  # Archivadas
      2.2*cm,  # Activas
      2.2*cm,  # Balance
      2.2*cm,  # Finalizadas
      2.2*cm,  # Canceladas
      2.2*cm,  # En plazo
      2.2*cm,  # Vencidas
      2.2*cm,  # Participación
      2.2*cm,  # Productividad
    ],
    repeatRows=1,
  )

  apply_table_style(team_table, len(team_rows), left_columns=2)

  story.append(team_table)

  story.append(Spacer(1, 0.8 * cm))

  doc.build(
    story,
    onFirstPage=draw_header,
    onLaterPages=draw_header,
    canvasmaker=NumberedCanvas,
  )

  return response
